"""Read a resume file and build a candidate Profile from it."""

from __future__ import annotations

import re
from datetime import date
from pathlib import Path

from .models import Profile
from .skills import canonicalize, extract_skills

# Seniority ladder shared with the matcher.
LEVELS = ["intern", "entry", "mid", "senior", "staff", "principal"]

MONTHS = {
    m: i + 1
    for i, names in enumerate(
        [("jan", "january"), ("feb", "february"), ("mar", "march"), ("apr", "april"), ("may",),
         ("jun", "june"), ("jul", "july"), ("aug", "august"), ("sep", "sept", "september"),
         ("oct", "october"), ("nov", "november"), ("dec", "december")]
    )
    for m in names
}

_MONTH_RE = r"(?:jan|feb|mar|apr|may|jun|jul|aug|sep|sept|oct|nov|dec)[a-z]*\.?"
_POINT = rf"(?:{_MONTH_RE}\s+\d{{4}}|\d{{1,2}}/\d{{4}}|\d{{4}})"
_END = rf"(?:{_POINT}|present|current|now|today)"
DATE_RANGE_RE = re.compile(rf"({_POINT})\s*(?:-|–|—|to|until)\s*({_END})", re.I)
YEARS_CLAIM_RE = re.compile(r"(\d{1,2})\+?\s*(?:years|yrs)\b(?:\s+of)?(?:\s+\w+){0,3}\s+experience", re.I)

TITLE_WORDS = (
    "engineer|developer|scientist|analyst|manager|designer|architect|consultant|administrator|"
    "specialist|researcher|lead|director|intern|programmer|technician|coordinator|associate|"
    "strategist|writer|recruiter|accountant|officer|sre|devops|assistant|representative|agent|advisor|"
    "executive|clerk|receptionist|teacher|tutor|nurse|editor|marketer|planner|supervisor"
)
# Up to 6 capitalised words (or "&", "and", "of") before a title word: "Travel Retail & Corporate Business APAC Intern".
# The prefix is case-sensitive so sentences like "Onboarded and trained an incoming intern" don't count.
TITLE_RE = re.compile(
    rf"^\s*((?:(?:[A-Z][\w/&+.()-]*|&|and|of|-|–)\s+){{0,6}}(?i:{TITLE_WORDS})\b"
    rf"(?:\s+(?:I|II|III|IV|[A-Z][\w/&+.-]*)){{0,3}}"
    rf"(?:\s+[–-]\s+[A-Z][\w/&+.-]*(?:\s+[A-Z][\w/&+.-]*){{0,3}})?)"  # "Senior Associate – UX Design"
)

# Resume section headings. Dates are only counted as work experience inside work sections,
# so degrees and side projects don't inflate your years.
SECTION_RE = re.compile(
    r"^\s*(?P<name>(?:professional |work |relevant |career )?(?:experience|employment(?: history)?|work history)|"
    r"education|academic\w*|projects?|personal projects|certifications?|skills|technical skills|core skills|"
    r"profile|summary|about me|awards|achievements|publications|volunteer\w*|interests|languages|references)"
    r"\s*:?\s*$",
    re.I,
)
WORK_SECTION_RE = re.compile(r"experience|employment|work history", re.I)


def read_resume_text(path: str | Path) -> str:
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(str(path))
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    if suffix == ".docx":
        import docx

        document = docx.Document(str(path))
        parts = [p.text for p in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        return "\n".join(parts)
    return path.read_text(encoding="utf-8", errors="ignore")


def _parse_point(token: str, *, is_end: bool) -> date | None:
    """Parse one side of a date range. End dates are exclusive (the month after), so
    "Jan 2020 - Dec 2020" and "2020 - 2020" both count as a full year."""
    token = token.strip().lower().rstrip(".")
    if token in {"present", "current", "now", "today"}:
        return date.today()
    year = month = None
    m = re.match(rf"({_MONTH_RE})\s+(\d{{4}})", token)
    if m:
        name = m.group(1).rstrip(".")
        year, month = int(m.group(2)), MONTHS.get(name) or MONTHS.get(name[:3]) or 1
    elif m := re.match(r"(\d{1,2})/(\d{4})", token):
        year, month = int(m.group(2)), max(1, min(12, int(m.group(1))))
    elif m := re.match(r"(\d{4})", token):
        year, month = int(m.group(1)), 12 if is_end else 1
    if year is None:
        return None
    if is_end:
        year, month = (year + 1, 1) if month == 12 else (year, month + 1)
    return date(year, month, 1)


def split_sections(text: str) -> dict[str, str]:
    """Split a resume into {heading: text}. Text before the first heading goes under "header"."""
    sections: dict[str, list[str]] = {"header": []}
    current = "header"
    for line in text.splitlines():
        m = SECTION_RE.match(line)
        if m:
            current = m.group("name").lower()
            sections.setdefault(current, [])
        else:
            sections[current].append(line)
    return {name: "\n".join(lines) for name, lines in sections.items()}


def work_text(text: str) -> str:
    """Just the work-experience part of a resume, or the whole resume if it has no such heading."""
    work = [body for name, body in split_sections(text).items() if WORK_SECTION_RE.search(name)]
    return "\n".join(work) if work else text


def estimate_years(text: str) -> float:
    """Estimate total professional experience by merging the date ranges in the work section(s).

    Overlapping jobs are counted once. Falls back to explicit claims like "7+ years of
    experience" when those are larger.
    """
    claims = [int(n) for n in YEARS_CLAIM_RE.findall(text)]
    text = work_text(text)
    intervals = []
    for start_s, end_s in DATE_RANGE_RE.findall(text):
        start, end = _parse_point(start_s, is_end=False), _parse_point(end_s, is_end=True)
        if start and end and 1970 <= start.year <= date.today().year and end >= start:
            intervals.append((start, min(end, date.today())))
    intervals.sort()
    merged: list[list[date]] = []
    for start, end in intervals:
        if merged and start <= merged[-1][1]:
            merged[-1][1] = max(merged[-1][1], end)
        else:
            merged.append([start, end])
    from_ranges = sum((e - s).days for s, e in merged) / 365.25
    return round(max([from_ranges] + claims), 1)


def extract_titles(text: str, limit: int = 6) -> list[str]:
    """Job titles from the work section, most recent first (as they appear in the resume)."""
    titles: list[str] = []
    for line in work_text(text).splitlines():
        line = re.sub(r"\s+", " ", line).strip()  # PDFs often pad columns with runs of spaces
        if not line or len(line) > 160 or line[0] in "•·-*–●▪◦":
            continue
        # Titles usually sit on their own line or before a separator such as "|", "," or "at".
        if "|" in line:
            head = line.split("|")[0]  # "Title | Company | Dates": a dash before the "|" is part of the title
        else:
            head = re.split(r"\s+(?:@|at|-|–|—)\s+|,|\t", line)[0]
        m = TITLE_RE.match(head)
        if m and len(m.group(1).split()) <= 7:
            title = re.sub(r"\s+", " ", m.group(1)).strip(" -,")
            if title.lower() not in (t.lower() for t in titles):
                titles.append(title)
        if len(titles) >= limit:
            break
    return titles


def level_for(years: float, titles: list[str]) -> int:
    if years < 2:
        level = 1
    elif years < 6:
        level = 2
    elif years < 10:
        level = 3
    else:
        level = 4
    # "Senior Associate" is a mid-level title at most firms, so it doesn't count as senior.
    joined = re.sub(r"senior associate", "associate", " ".join(titles).lower())
    if re.search(r"\bprincipal\b", joined):
        level = max(level, 5)
    elif re.search(r"\bstaff\b", joined):
        level = max(level, 4)
    elif re.search(r"\b(senior|sr\.?|lead)\b", joined):
        level = max(level, 3)
    if titles and all(re.search(r"\bintern", t, re.I) for t in titles) and years < 1:
        level = 0
    return level


def build_profile(text: str, extra_skills: list[str] | None = None, ignore_skills: list[str] | None = None,
                  titles: list[str] | None = None,
                  years_override: float | None = None) -> Profile:
    # extra_skills: ones the parser missed; ignore_skills: ones it picked up by mistake.
    skills = extract_skills(text) | {canonicalize(s) for s in (extra_skills or [])}
    skills -= {canonicalize(s) for s in (ignore_skills or [])}
    found_titles = extract_titles(text)
    years = years_override if years_override is not None else estimate_years(text)
    management = bool(re.search(r"\b(manager|director|head of|vp|vice president)\b", " ".join(found_titles), re.I))
    return Profile(
        text=text,
        skills=skills,
        titles=titles or found_titles,
        years_experience=years,
        level=level_for(years, found_titles),
        has_management=management,
    )


def load_profile(path: str | Path, **kwargs) -> Profile:
    return build_profile(read_resume_text(path), **kwargs)
